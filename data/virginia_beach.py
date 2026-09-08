"""Loader for the Virginia Beach OHCA replication instance."""

from pathlib import Path

import numpy as np


EARTH_RADIUS_KM = 6371.0088


def _read_coordinates(workbook_path: Path, sheet_name: str, latitude: str, longitude: str):
    """Return latitude/longitude coordinates from a named workbook sheet."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - exercised by users without extras
        raise RuntimeError(
            "Virginia Beach replication requires openpyxl. "
            "Install the dependencies in requirements.txt."
        ) from exc

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        columns = {name: index for index, name in enumerate(headers)}
        missing = [name for name in (latitude, longitude) if name not in columns]
        if missing:
            raise ValueError(f"Missing columns in {sheet_name}: {', '.join(missing)}")

        coordinates = [
            (float(row[columns[latitude]]), float(row[columns[longitude]]))
            for row in sheet.iter_rows(min_row=2, values_only=True)
            if row[columns[latitude]] is not None and row[columns[longitude]] is not None
        ]
    finally:
        workbook.close()

    if not coordinates:
        raise ValueError(f"No valid coordinates found in {sheet_name}.")
    return np.asarray(coordinates, dtype=float)


def _haversine_km(facilities: np.ndarray, demands: np.ndarray) -> np.ndarray:
    """Compute all facility-demand great-circle distances in kilometres."""
    facility_latitude = np.deg2rad(facilities[:, 0])[:, np.newaxis]
    facility_longitude = np.deg2rad(facilities[:, 1])[:, np.newaxis]
    demand_latitude = np.deg2rad(demands[:, 0])[np.newaxis, :]
    demand_longitude = np.deg2rad(demands[:, 1])[np.newaxis, :]

    delta_latitude = demand_latitude - facility_latitude
    delta_longitude = demand_longitude - facility_longitude
    a = (
        np.sin(delta_latitude / 2.0) ** 2
        + np.cos(facility_latitude)
        * np.cos(demand_latitude)
        * np.sin(delta_longitude / 2.0) ** 2
    )
    return 2.0 * EARTH_RADIUS_KM * np.arcsin(np.sqrt(a))


def load_virginia_beach_instance(
    workbook_path: str | Path,
    coverage_radius: float = 4.0,
    max_coverage_radius: float = 8.0,
    decay_rate: float = 0.5,
    facility_budget: int = 10,
):
    """Build the MMCLBP instance used in the Virginia Beach case study.

    The workbook contains 40 base stations and 2,706 OHCA events. Distances are
    recomputed with the Haversine formula, as specified in the manuscript.
    """
    from .data import Data

    workbook_path = Path(workbook_path)
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Virginia Beach workbook not found: {workbook_path}")

    facilities = _read_coordinates(workbook_path, "Base_Stations", "Latitude", "Longitude")
    demands = _read_coordinates(workbook_path, "OHCAs", "Latitude", "Longitude")

    data = Data()
    data.I = list(range(len(facilities)))
    data.J = list(range(len(demands)))
    data.facility_coord = facilities
    data.demand_coord = demands
    data.coverage_radius = coverage_radius
    data.max_coverage_radius = max_coverage_radius
    data.decay_rate = decay_rate
    data.facility_cost = np.ones(len(data.I), dtype=int)
    data.blocker_cost = np.ones(len(data.I), dtype=int)
    data.facility_budget = facility_budget
    data.demand_weight = np.ones(len(data.J), dtype=int)
    data.target_coverage = 0.0
    data.distance = _haversine_km(facilities, demands)
    data._compute_coverage_matrix()
    data._build_cover_sets()
    data.bigM = np.asarray(data.prob @ data.demand_weight, dtype=np.float64)
    return data
